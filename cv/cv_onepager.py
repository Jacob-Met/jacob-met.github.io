#!/usr/bin/env python3
"""Deterministic one-page CV generator (PDF + DOCX) from the public record source/content.json.

Same input -> byte-identical output. Public records only: the record must say visibility=public,
every claim needs a source link, and every link must be https on the site's approved host set
(same rules as source/build.py). Nothing is fetched and nothing is published; it only writes files.

    python cv/cv_onepager.py --out cv-out            # writes Jacob-Metoyer-CV.pdf / .docx
    python cv/cv_onepager.py --content other.json --out /tmp/x

Needs reportlab and python-docx (pip install reportlab python-docx). The site build itself stays
stdlib-only; this tool lives outside source/ so the site CI is unaffected.
"""
from __future__ import annotations
import argparse, hashlib, io, json, sys, zipfile
from pathlib import Path
from urllib.parse import urlsplit

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_CONTENT = ROOT / 'source' / 'content.json'
STEM = 'Jacob-Metoyer-CV'
ALLOWED_HOSTS = {'github.com', 'www.linkedin.com', 'www.csulb.edu', 'www.joshilab.org'}
FIXED_ZIP_TIME = (2026, 1, 1, 0, 0, 0)


class RecordError(ValueError):
    pass


# ---------------------------------------------------------------- record -> neutral model
def _check_url(u: str) -> str:
    p = urlsplit(u)
    if p.scheme != 'https' or p.hostname not in ALLOWED_HOSTS or p.username or p.password or p.port:
        raise RecordError(f'link outside approved public set: {u}')
    return u


def _need_sources(item: dict, what: str) -> None:
    if not item.get('sources'):
        raise RecordError(f'claim without a source link: {what}')
    for s in item['sources']:
        _check_url(s['url'])


def _basis(b: str) -> str:
    return 'self-reported' if b.startswith('self-reported') else ''


def load(path: Path = DEFAULT_CONTENT) -> dict:
    return json.loads(Path(path).read_text(encoding='utf-8'))


def model(rec: dict) -> dict:
    """Neutral, ordered content model: header, then sections of (heading, [entry...]).
    entry = dict(title, meta, body, tag, link)."""
    if rec.get('visibility') != 'public':
        raise RecordError('record is not marked public')
    for i in rec.get('identities', []):
        _check_url(i['url'])
    edu = rec['education']; _need_sources(edu, 'education')
    sections = []
    sections.append(('Education', [dict(
        title=edu['institution'], meta=edu['program'], body=edu['degrees'], tag=_basis(edu['basis']), link='')]))
    sections.append(('Research experience', []))
    for r in rec['research']:
        _need_sources(r, r['title'])
        sections[-1][1].append(dict(title=r['title'], meta=f"{r['org']} · {r['period']}", body=r['role'],
                                    tag=_basis(r['basis']), link=''))
    sections.append(('Projects', []))
    for p in rec['projects']:
        arts = p['artifacts']
        if not arts: raise RecordError(f"project without artifact link: {p['title']}")
        for a in arts: _check_url(a['url'])
        sections[-1][1].append(dict(title=p['title'], meta=f"{p['kind']} · {p['facts']}", body=p['summary'],
                                    tag='', link=arts[0]['url']))
    ind = rec.get('independent')
    if ind:
        _need_sources(ind, 'independent research')
        sections.append(('Independent work', [dict(title='Human-AI collaboration and agentic systems', meta='',
                                                   body=ind['text'], tag=_basis(ind['basis']), link='')]))
    w = rec.get('verified_writing') or []
    if w:
        for x in w:
            for l in x['links']: _check_url(l['url'])
        sections.append(('Technical writing', [dict(title=x['title'], meta=f"{x['venue']} · {x['date']}", body='',
                                                     tag='', link=x['links'][0]['url']) for x in w]))
    return dict(name=rec['name'], headline=rec['headline'], summary=rec['summary'][0],
                identities=[(i['label'], i['url']) for i in rec.get('identities', [])],
                updated=rec['updated'], sections=sections,
                note='Entries marked self-reported are drawn from my own account; public repositories and '
                     'dated snapshots document the software work. Record dated ' + rec['updated'] + '.')


# ---------------------------------------------------------------- PDF (reportlab)
def render_pdf(m: dict) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.colors import HexColor
    from reportlab.lib.enums import TA_LEFT
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    from xml.sax.saxutils import escape

    ink, mute, acc = HexColor('#1a1a1a'), HexColor('#555555'), HexColor('#1f4e79')
    base = dict(fontName='Helvetica', textColor=ink, alignment=TA_LEFT)
    s_name = ParagraphStyle('n', fontName='Helvetica-Bold', fontSize=20, leading=23, textColor=ink)
    s_head = ParagraphStyle('h', fontSize=9, leading=11.5, **{**base, 'textColor': mute})
    s_sec = ParagraphStyle('s', fontName='Helvetica-Bold', fontSize=8.5, leading=10, textColor=acc, spaceBefore=5, spaceAfter=1)
    s_t = ParagraphStyle('t', fontSize=8.6, leading=10.4, **base)
    s_b = ParagraphStyle('b', fontSize=7.9, leading=9.6, leftIndent=0, **base)
    s_n = ParagraphStyle('nt', fontSize=6.8, leading=8.2, **{**base, 'textColor': mute})

    def esc(t): return escape(t, {'"': '&quot;'})
    story = [Paragraph(esc(m['name']), s_name), Paragraph(esc(m['headline']), s_head),
             Paragraph('  ·  '.join(f'<a href="{esc(u)}" color="#1f4e79">{esc(l)}: {esc(u.split("//", 1)[1])}</a>'
                                    for l, u in m['identities']), s_head),
             Spacer(1, 2), Paragraph(esc(m['summary']), s_b)]
    for heading, entries in m['sections']:
        story += [Paragraph(heading.upper(), s_sec), HRFlowable(width='100%', thickness=0.4, color=mute, spaceAfter=2)]
        for e in entries:
            head = f"<b>{esc(e['title'])}</b>"
            if e['meta']: head += f" <font color='#555555'>— {esc(e['meta'])}</font>"
            if e['tag']: head += f" <font color='#1f4e79' size='6.6'>[{esc(e['tag'])}]</font>"
            if e['link']: head += f" <a href='{esc(e['link'])}' color='#1f4e79'><font size='6.6'>{esc(e['link'].split('//', 1)[1])}</font></a>"
            story.append(Paragraph(head, s_t))
            if e['body']: story.append(Paragraph(esc(e['body']), s_b))
            story.append(Spacer(1, 1.6))
    story += [Spacer(1, 3), Paragraph(esc(m['note']), s_n)]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=letter, leftMargin=40, rightMargin=40, topMargin=34, bottomMargin=30,
                            title=f"{m['name']} — CV", author=m['name'], subject='One-page CV', creator='cv_onepager.py',
                            invariant=1)  # invariant: fixed dates + document ID => reproducible bytes
    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------- DOCX (python-docx, zip normalised)
def render_docx(m: dict) -> bytes:
    import datetime
    from docx import Document
    from docx.shared import Pt, Inches, RGBColor
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.6); sec.top_margin = Inches(0.5); sec.bottom_margin = Inches(0.45)
    st = d.styles['Normal']; st.font.name = 'Liberation Sans'; st.font.size = Pt(8.5)
    st.paragraph_format.space_after = Pt(0); st.paragraph_format.space_before = Pt(0)
    mute, acc = RGBColor(0x55, 0x55, 0x55), RGBColor(0x1f, 0x4e, 0x79)

    def para(text='', size=None, bold=False, color=None, before=0, after=0):
        p = d.add_paragraph(); p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
        if text:
            r = p.add_run(text); r.bold = bold
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = color
        return p

    def run(p, text, size=None, bold=False, color=None):
        r = p.add_run(text); r.bold = bold
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        return r

    para(m['name'], 20, True)
    para(m['headline'], 9, color=mute)
    para('  ·  '.join(f'{l}: {u.split("//", 1)[1]}' for l, u in m['identities']), 9, color=mute, after=2)
    para(m['summary'], 8.5, after=1)
    for heading, entries in m['sections']:
        p = para(heading.upper(), 8.5, True, acc, before=5, after=1)
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        pPr = p._p.get_or_add_pPr(); bd = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
        for k, v in (('w:val', 'single'), ('w:sz', '4'), ('w:space', '1'), ('w:color', '555555')): b.set(qn(k), v)
        bd.append(b); pPr.append(bd)
        for e in entries:
            p = para(before=1)
            run(p, e['title'], 8.8, True)
            if e['meta']: run(p, ' — ' + e['meta'], 8.8, False, mute)
            if e['tag']: run(p, f" [{e['tag']}]", 7, False, acc)
            if e['link']: run(p, ' ' + e['link'].split('//', 1)[1], 7, False, acc)
            if e['body']: para(e['body'], 8.2)
    para(m['note'], 7, color=mute, before=4)

    cp = d.core_properties
    fixed = datetime.datetime(2026, 1, 1, 0, 0, 0)
    cp.author = m['name']; cp.last_modified_by = m['name']; cp.title = f"{m['name']} — CV"
    cp.created = fixed; cp.modified = fixed; cp.revision = 1; cp.comments = ''
    raw = io.BytesIO(); d.save(raw)
    # Normalise the zip: sorted entries, fixed timestamps, fixed attrs => reproducible bytes.
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(raw.getvalue())) as zin, zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as zout:
        for name in sorted(zin.namelist(), key=lambda n: (n != '[Content_Types].xml', n)):
            zi = zipfile.ZipInfo(name, FIXED_ZIP_TIME); zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o600 << 16; zi.create_system = 3
            zout.writestr(zi, zin.read(name))
    return out.getvalue()


# ---------------------------------------------------------------- driver
def generate(content: Path, out: Path) -> dict:
    m = model(load(content))
    out.mkdir(parents=True, exist_ok=True)
    res = {}
    for ext, fn in (('pdf', render_pdf), ('docx', render_docx)):
        data = fn(m); p = out / f'{STEM}.{ext}'
        tmp = p.with_suffix(p.suffix + '.tmp'); tmp.write_bytes(data); tmp.replace(p)
        res[p.name] = hashlib.sha256(data).hexdigest()
    (out / f'{STEM}.sha256.json').write_text(json.dumps(res, indent=1, sort_keys=True) + '\n', encoding='utf-8')
    return res


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--content', type=Path, default=DEFAULT_CONTENT)
    ap.add_argument('--out', type=Path, default=Path('cv-out'))
    a = ap.parse_args(argv)
    try:
        res = generate(a.content, a.out)
    except RecordError as e:
        print(f'REFUSED: {e}', file=sys.stderr); return 2
    for k, v in res.items(): print(f'{v}  {a.out / k}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
